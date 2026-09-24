"""Gera as planilhas da rodada 22/09: uma por editor + a de áudio/roteiro do Tiago.

Entradas: plano-consolidado.csv (consolidar.py) e consulta-editores_22-09-2026.csv
(export da aba Setembro/2026 com editor, links e estado de AF/AD).
"""
import csv
from pathlib import Path
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.hyperlink import Hyperlink

AQUI = Path(__file__).parent
PASTA_MAE = 'https://drive.google.com/drive/u/3/folders/18UAhFf13x10S9EhHbFQOzrpDY8VrScj_'

HEAD_FILL = PatternFill('solid', fgColor='FF1F4E4A')
HEAD_FONT = Font(b=True, color='FFFFFFFF')
WRAP = Alignment(wrap_text=True, vertical='top')


def norm(s):
    return s.strip().lstrip('0')


def nome_editor(e):
    e = e.strip()
    return {'Bidóia': 'Bidoia', '': 'SEM-EDITOR'}.get(e, e)


plano = [r for r in csv.DictReader(open(AQUI / 'plano-consolidado.csv', encoding='utf-8-sig')) if r['em_aberto']]
sheet = {norm(r['sku']): r for r in csv.DictReader(open(AQUI / 'consulta-editores_22-09-2026.csv', encoding='utf-8-sig'))}

ordem_d = {d: i for i, d in enumerate(['08.09', '10.09', '11.09', '14.09', '15.09', '16.09', '17.09', '18.09', '21.09'])}
ordem_af = {'REINSERÇÃO CENAS IA': 0, 'REINSERÇÃO CENAS IA/áudio': 1, 'REINSERÇÃO CENAS IA/Alteração': 2,
            'Alteração não feita': 3, 'Alteração': 4, 'Alteração/áudio': 5, 'Alteração áudio': 6}

itens = []
for r in plano:
    s = sheet.get(norm(r['sku']), {})
    itens.append(dict(
        diaria=r['diaria'], cliente=r['cliente'], sku=r['sku'], produto=r['produto'],
        tipo=s.get('tipo', ''), af=r['af'], nota=r['nota'], obs=r['obs'], audio=r['audio'],
        vai_editor=r['vai_editor'], vai_audio=r['vai_audio'],
        editor=nome_editor(s.get('editor', '')) if s else 'SEM-EDITOR',
        link_drive=s.get('link_drive', ''), link_roteiro=s.get('link_roteiro', ''),
        linha=s.get('linha', ''), af_planilha=s.get('af', ''), ad=s.get('ad_enviado_cliente', ''),
    ))
itens.sort(key=lambda x: (ordem_d[x['diaria']], ordem_af.get(x['af'], 9), x['sku']))


def planilha(path, titulo, sub, cab, larg, linhas, link_cols=()):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Ajustes'
    ws['A1'] = titulo
    ws['A1'].font = Font(sz=14, b=True)
    ws['A2'] = sub
    ws['A3'] = f'Pasta-mãe no Drive: {PASTA_MAE}'
    for i, (h, w) in enumerate(zip(cab, larg), 1):
        c = ws.cell(row=5, column=i, value=h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        ws.column_dimensions[c.column_letter].width = w
    for r, vals in enumerate(linhas, 6):
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment = WRAP
            if i in link_cols and v:
                c.hyperlink = v
                c.value = 'abrir'
                c.font = Font(color='FF0563C1', u='single')
    ws.freeze_panes = 'A6'
    ws.auto_filter.ref = f'A5:{ws.cell(row=5, column=len(cab)).column_letter}{max(5, 5 + len(linhas))}'
    wb.save(path)


def o_que_fazer(x):
    t = x['af']
    if x['nota']:
        t += ' — ' + x['nota']
    if x['af_planilha'] == 'Aprovado':
        t += ' — VERIFICAR: na planilha consta Aprovado, mas o cliente pediu esta alteração'
    if x['editor'] == 'SEM-EDITOR':
        t += ' — sem editor na planilha' + ('' if x['linha'] else ' (SKU não encontrado na aba Setembro/2026)')
    return t


# --- uma planilha por editor
por_editor = defaultdict(list)
for x in itens:
    if x['vai_editor']:
        por_editor[x['editor']].append(x)

cab = ['Diária', 'SKU', 'Produto', 'Tipo', 'O que fazer', 'Observação do cliente', 'Pasta']
larg = [9, 12, 42, 10, 30, 66, 8]
resumo = []
for ed, xs in sorted(por_editor.items()):
    ds = sorted({x['diaria'] for x in xs}, key=ordem_d.get)
    n_r = sum(1 for x in xs if x['af'].startswith('REINSERÇÃO'))
    sub = f'Diárias {", ".join(ds)} — {len(xs)} vídeo(s): {n_r} com reinserção de cenas IA, {len(xs) - n_r} de edição'
    planilha(AQUI / f'ajustes_{ed}.xlsx', f'Ajustes do cliente — {ed}', sub, cab, larg,
             [[x['diaria'], x['sku'], x['produto'], x['tipo'], o_que_fazer(x), x['obs'], x['link_drive']] for x in xs],
             link_cols=(7,))
    resumo.append((ed, len(xs), n_r))

# --- áudio / roteiro (Tiago)
aud = [x for x in itens if x['vai_audio']]
planilha(AQUI / 'audio-e-roteiro_CONTROLE.xlsx', 'Alterações de áudio e roteiro — pedir regravação',
         f'{len(aud)} vídeo(s). "VERIFICAR" = conferir no vídeo se o erro é da locução ou do lettering antes de pedir.',
         ['Diária', 'SKU', 'Produto', 'O que pedir', 'Também vai ao editor?', 'Editor', 'Observação do cliente', 'Roteiro', 'Pasta'],
         [9, 12, 38, 50, 12, 11, 60, 8, 8],
         [[x['diaria'], x['sku'], x['produto'], x['audio'], 'sim' if x['vai_editor'] else 'não', x['editor'], x['obs'],
           x['link_roteiro'], x['link_drive']] for x in aud],
         link_cols=(8, 9))

for ed, n, n_r in resumo:
    print(f'{ed:12} {n:3}  (reinserção {n_r})')
print('áudio', len(aud))

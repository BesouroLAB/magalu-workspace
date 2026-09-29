"""Rodada 25/09 — diárias 21.09 (Dany) e 23.09 (Kauê).

Gera uma planilha por editor + a de áudio/roteiro do Tiago.
Editor e links vêm do export da aba Setembro/2026 (setembro-2026_export-25.09.csv).
"""
import csv
from pathlib import Path
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

AQUI = Path(__file__).parent
PASTA_MAE = 'https://drive.google.com/drive/u/3/folders/18UAhFf13x10S9EhHbFQOzrpDY8VrScj_'
HEAD_FILL = PatternFill('solid', fgColor='FF1F4E4A')
HEAD_FONT = Font(b=True, color='FFFFFFFF')
WRAP = Alignment(wrap_text=True, vertical='top')

R = 'REINSERÇÃO CENAS IA'
RA = 'REINSERÇÃO CENAS IA/Alteração'
RAU = 'REINSERÇÃO CENAS IA/áudio'
A = 'Alteração'

# (diária, AD esperado, sku, produto, af, nota, áudio, obs do cliente)
ITENS = [
    ('21.09', '21/09/2026', '240212900', 'Triciclo Infantil de Pedal Mototico Bandeirante', R, '', '',
     '0:07 - O triciclo é da Bandeirante, mas aparece o nome da marca concorrente: Magic Toys'),
    ('21.09', '21/09/2026', '240367300', 'Conjunto para Servir Oxford Brisa 2 Peças de Cerâmica', R, '', '',
     'IA alucinou o produto no vídeo todo'),
    ('21.09', '21/09/2026', '240187100', 'Pista Marvel Spider-Man Free Fall Candide', R, '', '',
     '0:02 - A pista ta diferente\n0:09 - A pista e o carrinho tão diferentes\n0:12 - A pista ta diferente\n0:18 - A pista ta diferente'),
    ('21.09', '21/09/2026', '240554600', 'Câmera Inteligente Wi-Fi ELG Full HD Full Color SHCF602', RAU, '',
     'HECTOR (troca de trecho) — VERIFICAR: 0:11 "fala" que tem proteção contra chuva, sem respaldo do fabricante. Se for locução, tirar o trecho; se for lettering, é do editor',
     '0:10 - Na ficha técnica diz para evitar instalar em ambientes com umidade. Na imagem, aparece a câmera toda molhada.\n0:11 - Fala que ela tem proteção contra chuva, mas no site do fabricante não fala nada a respeito.\n0:13 - Do jeito que a câmera está, parece que ela tem luz que ilumina o ambiente. Mas não temos nenhuma informação a respeito.'),
    ('21.09', '21/09/2026', '240411200', 'Andador Infantil Macaco Mattel Fisher-Price Aprenda Comigo Musical', R, '', '',
     '0:04 - O produto tá diferente'),
    ('21.09', '21/09/2026', '238830500', 'Estante Multiuso 5 Prateleiras Multivisão Stronger', A, 'incluir TL', '',
     '0:20 - Incluir TL: Imagem de sugestão de uso. Verifique as características do produto.'),
    ('21.09', '21/09/2026', '240554800', 'Fechadura Digital ELG SHFD701 com Senha e Cartão de Embutir Wi-Fi', R, '', '',
     '0:07 - Aqui ela foi instalada numa porta sem maçaneta'),
    ('21.09', '21/09/2026', '241007800', 'LEGO City Assalto ao Trem da Polícia 60508', R, '', '',
     '0:03 - Tem dois carrinhos vermelhos diferentes do carrinho que vem no kit.\n0:17 - Tem dois carrinhos vermelhos diferentes do carrinho que vem no kit.'),
    ('21.09', '21/09/2026', '241501800', 'Guarda-roupa Casal com Espelho e Sapateiro 4 Gavetas 2 Portas de Correr Conquista Móveis Prime Chicago', R, '', '',
     '0:13 - Deu um bug na imagem, em cima de uma das portas do guarda-roupas.'),
    ('21.09', '21/09/2026', '241004500', 'LEGO Star Wars Droide Astromecânico BB-8', R, '', '',
     '0:16 - a porta redonda da frente tá diferente'),
    ('21.09', '21/09/2026', '240996700', 'LEGO City Caminhão Betoneira 60478 371 Peças', R, '', '',
     '0:03 - O LEGO vem com 03 bonecos, mas estão aparecendo 05\n0:17 - O LEGO vem com 03 bonecos, mas estão aparecendo 05'),
    ('21.09', '21/09/2026', '240264500', 'LEGO Cabana da Vila Fusha 75636', R, '', '',
     '0:03 - Alucinou a roupa da boneca e do boneco de chapéu\n0:17 - Alucinou a roupa da boneca e do boneco de chapéu'),
    ('21.09', '21/09/2026', '240101600', 'Guarda-roupa Solteiro com Espelho 4 Gavetas 2 Portas Santos Andirá Democrata', A,
     'régua de medidas + mostrar as duas opções de porta (flex)', '',
     '0:07- Precisa desenhar a régua de medidas\nO produto é flex, tem duas opções de porta. Mas isso não é mencionado no vídeo.'),
    ('21.09', '21/09/2026', '241001400', 'LEGO Minecraft Batalha Wither 21590 494 Peças', R, '', '',
     '0:03 - Tem um foguete do lado esquerdo, saindo do bloco vermelho. Não faz parte do kit.\n0:17 - Tem um foguete do lado esquerdo, saindo do bloco vermelho. Não faz parte do kit.'),
    ('21.09', '21/09/2026', '240881700', 'Caixa de Som LG Xboom by Will.i.am Bounce Bluetooth Amplificada Portátil 40W', R, '', '',
     '0:15 - A caixa tá diferente'),
    ('21.09', '21/09/2026', '228760200', 'Headset Gamer Philco PHS11V PC 5.1 P2', R, '', '',
     '0:16 - IA alucinou o cabo'),
    ('21.09', '21/09/2026', '240350000', 'Dinossauro de Brinquedo ML102 Emite Som', R, '', '',
     '0:04 - Tem uns buracos na perna dele, mas o de verdade não tem\n0:11- Aqui ele abre a boca, mas o de verdade não abre.'),
    ('21.09', '21/09/2026', '240347600', 'Cozinha Infantil Kit Almoço Lulie Kids', R, '', '',
     '0:03 - A colher amarela tá duplicada.\n0:20 - A colher amarela tá duplicada'),
    ('21.09', '21/09/2026', '220722300', 'Remo Kikos 218CA 12 Níveis de Regulagem', R, '', '',
     '0:03 - IA alucinou o produto\n0:10 - IA alucinou o produto\n0:15 - IA alucinou o produto'),
    ('21.09', '21/09/2026', '241002600', 'LEGO Minecraft Nether e o Portal do End 21584', R, '', '',
     '0:06 - IA Alucinou completamente o brinquedo\n0:08 - IA Alucinou completamente o brinquedo\n0:16 - IA Alucinou completamente o brinquedo'),
    ('21.09', '21/09/2026', '235626300', 'Brinquedo Educativo Tópi Bombeirinho', R, '', '',
     '0:03 - A peça amarela tá duplicada.\n0:10 - A peça amarela tá duplicada.\n0:16 - Parte de dentro tá diferente.\n0:20 - A peça amarela tá duplicada.'),
    ('21.09', '21/09/2026', '240993900', 'LEGO City Transportador de Motocicletas 60491', R, '', '',
     '0:06 - A IA alucinou a roupa do boneco'),
    ('21.09', '21/09/2026', '240999000', 'LEGO Disney Frozen Castelo Elsa Passeio de Neve', R, '', '',
     '0:06 - O segundo andar do castelo tá diferente\n0:10 - O cabelo da elsa e o vestido estão diferentes\n0:20 - Tem 2 olafs'),
    ('21.09', '21/09/2026', '204592700', 'Kit Churrasco 12 Peças', R, '', '',
     '0:14 - Tá escrito "churrasco" na faca'),
    ('21.09', '21/09/2026', '234981600', 'Frigideira Antiaderente Tramontina de Alumínio Paris Vermelha 24cm com Espátula', A,
     'áudio errado: está com a locução de um berço portátil — trocar pelo áudio da frigideira (Roteiro_e_Audio da pasta) + tirar a linha preta', '',
     '0:01 - Aparece uma linha preta na lateral esquerda\nO vídeo é uma frigideira, o áudio tá de um berço portátil'),
    ('21.09', '21/09/2026', '240009100', 'Painel para TV até 75” Caemmun Dal', RA, '', '',
     '0:07 - Maçaneta diferente\n0:10 - Produto tá diferente\n0:18 - Do jeito que as opções de cores foram mostradas não fica nítido.\nFaltou régua de medidas'),
    ('23.09', '23/09/2026', '241556800', 'Cama de Solteiro Conquista Móveis Rubi Branca', R, '', '',
     '00:15 – A cabeceira aparece vazada, porém, no produto original, ela é fechada. Além disso, ao final da cena, a IA adiciona um elemento voando.\n00:20 – A imagem da cama aparece achatada, comprometendo a proporção e a percepção das dimensões reais do produto.'),
    ('23.09', '23/09/2026', '240995800', 'LEGO City Helicóptero da Guarda Costeira 60503', R, '', '',
     '00:04 – A hélice do helicóptero aparece girando sozinha, sem que seja demonstrada nenhuma ação que provoque o movimento. Isso pode gerar uma interpretação incorreta sobre o funcionamento do produto.'),
    ('23.09', '23/09/2026', '241004400', 'LEGO Speed Champions Bugatti Vision Gran Turismo', A, 'estender a cena até o fim da locução', '',
     'A cena termina antes do final da locução, fazendo com que parte da fala seja cortada.'),
    ('23.09', '23/09/2026', '240792400', 'Colchão Solteiro Umaflex de Espuma D33 20x88x188cm Parma', R, '', '',
     '00:17 – O colchão aparece muito maior do que o modelo original, dando a impressão de que possui um box integrado.'),
    ('23.09', '23/09/2026', '238838800', 'Sofá 3 Lugares Suede Adágio Linoforte', A, 'régua de medidas', '',
     'Na cena em que são apresentadas as medidas, apenas a régua do comprimeto está representada corretamente, enquanto os demais números aparecem “soltos”.'),
    ('23.09', '23/09/2026', '240138600', 'Jogo de Panelas Neuhaus Revestimento Cerâmico de Alumínio Forjado Preto Fundo Triplo 8 Peças Ceramic Prime', A,
     'reduzir e variar as cenas', '',
     'Primeira e última cenas – As cenas são praticamente iguais e permanecem muito tempo em tela, deixando o vídeo repetitivo e monótono.\nCena do jogo completo – A cena inicia com movimento, mas logo em seguida fica totalmente estática, prejudicando o ritmo do vídeo.\nReduzir a duração das cenas e variar a apresentação para tornar o vídeo mais dinâmico'),
]


def norm(s):
    s = s.strip()
    if 'E' in s.upper():
        try:
            s = str(round(float(s.replace(',', '.'))))
        except ValueError:
            pass
    return s.split('.')[0].lstrip('0')


linhas = defaultdict(list)
with open(AQUI / 'setembro-2026_export-25.09.csv', encoding='utf-8-sig') as fh:
    rd = csv.reader(fh)
    H = [h.strip() for h in next(rd)]
    ix = {n: H.index(n) for n in ['Sku', 'Tipo', 'Link Roteiro', 'Link Drive', 'Editor', 'Data/Enviado/Cliente', 'Aprov Vídeo']}
    for n, row in enumerate(rd, 2):
        g = lambda k: row[ix[k]].strip() if ix[k] < len(row) else ''
        linhas[norm(g('Sku'))].append(dict(linha=n, tipo=g('Tipo'), roteiro=g('Link Roteiro'), drive=g('Link Drive'),
                                           editor=g('Editor'), ad=g('Data/Enviado/Cliente'), af=g('Aprov Vídeo')))

itens, avisos = [], []
for d, ad, sku, prod, af, nota, audio, obs in ITENS:
    cand = linhas.get(norm(sku), [])
    certa = [c for c in cand if c['ad'] == ad] or cand
    s = certa[-1] if certa else {}
    if not cand:
        avisos.append(f'{sku} {prod[:40]}: SKU não encontrado na aba Setembro/2026')
    elif not [c for c in cand if c['ad'] == ad]:
        avisos.append(f'{sku} {prod[:40]}: nenhuma linha com AD {ad} (usada a linha {s["linha"]}, AD {s["ad"]})')
    ed = {'Bidóia': 'Bidoia'}.get(s.get('editor', ''), s.get('editor', '')) or 'SEM-EDITOR'
    extra = nota
    if s.get('af') == 'Aprovado':
        extra += (' — ' if extra else '') + 'VERIFICAR: na planilha consta Aprovado, mas o cliente pediu esta alteração'
    if ed == 'SEM-EDITOR':
        extra += (' — ' if extra else '') + ('sem editor na planilha' if s else 'SKU não encontrado na aba Setembro/2026')
    itens.append(dict(d=d, sku=sku, prod=prod, af=af, extra=extra, audio=audio, obs=obs, ed=ed, **{k: s.get(k, '') for k in ('tipo', 'roteiro', 'drive')}))


def planilha(path, titulo, sub, cab, larg, rows, link_cols=()):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Ajustes'
    ws['A1'] = titulo
    ws['A1'].font = Font(sz=14, b=True)
    ws['A2'] = sub
    ws['A3'] = f'Pasta-mãe no Drive: {PASTA_MAE}'
    for i, (h, w) in enumerate(zip(cab, larg), 1):
        c = ws.cell(row=5, column=i, value=h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        ws.column_dimensions[c.column_letter].width = w
    for r, vals in enumerate(rows, 6):
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment = WRAP
            if i in link_cols and v:
                c.hyperlink, c.value = v, 'abrir'
                c.font = Font(color='FF0563C1', u='single')
    ws.freeze_panes = 'A6'
    ws.auto_filter.ref = f'A5:{ws.cell(row=5, column=len(cab)).column_letter}{5 + len(rows)}'
    wb.save(path)


por_ed = defaultdict(list)
for x in itens:
    por_ed[x['ed']].append(x)
for ed, xs in sorted(por_ed.items()):
    n_r = sum(x['af'].startswith('REINSERÇÃO') for x in xs)
    planilha(AQUI / f'ajustes_{ed}.xlsx', f'Ajustes do cliente — {ed}',
             f'Diárias {", ".join(sorted({x["d"] for x in xs}))} — {len(xs)} vídeo(s): {n_r} com reinserção de cenas IA, {len(xs) - n_r} de edição',
             ['Diária', 'SKU', 'Produto', 'Tipo', 'O que fazer', 'Observação do cliente', 'Pasta'], [9, 12, 42, 10, 30, 66, 8],
             [[x['d'], x['sku'], x['prod'], x['tipo'], x['af'] + (' — ' + x['extra'] if x['extra'] else ''), x['obs'], x['drive']] for x in xs],
             link_cols=(7,))
    print(f'{ed:12} {len(xs):3}  (reinserção {n_r})')

aud = [x for x in itens if x['audio']]
planilha(AQUI / 'audio-e-roteiro_CONTROLE.xlsx', 'Alterações de áudio e roteiro — pedir regravação',
         f'{len(aud)} vídeo(s). Pronúncia → Drika; corte/troca de trecho → Hector. "VERIFICAR" = conferir se é locução ou lettering.',
         ['Diária', 'SKU', 'Produto', 'O que pedir', 'Editor', 'Observação do cliente', 'Roteiro', 'Pasta'],
         [9, 12, 38, 50, 11, 60, 8, 8],
         [[x['d'], x['sku'], x['prod'], x['audio'], x['ed'], x['obs'], x['roteiro'], x['drive']] for x in aud],
         link_cols=(7, 8))
print('áudio', len(aud))
for a in avisos:
    print('AVISO', a)

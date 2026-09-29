"""Rodada 28/09 — diária 24.09 (Kauê).

Gera uma planilha por editor + a de áudio/roteiro do Tiago.
Editor e links vêm do export da aba Setembro/2026 (setembro-2026_export.csv).
"""
import csv
from pathlib import Path
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

AQUI = Path(__file__).parent
PASTA_MAE = 'https://drive.google.com/drive/u/3/folders/18UAhFf13x10S9EhHbFQOzrpDY8VrScj_'
HEAD_FILL = PatternFill('solid', fgColor='FF1F4E4A')
HEAD_FONT = Font(b=True, color='FFFFFFFF')
WRAP = Alignment(wrap_text=True, vertical='top')

from itens import *


def norm(s):
    s = s.strip()
    if 'E' in s.upper():
        try:
            s = str(round(float(s.replace(',', '.'))))
        except ValueError:
            pass
    return s.split('.')[0].lstrip('0')


linhas = defaultdict(list)
with open(AQUI / 'setembro-2026_export.csv', encoding='utf-8-sig') as fh:
    rd = csv.reader(fh)
    H = [h.strip() for h in next(rd)]
    ix = {n: H.index(n) for n in ['Sku', 'Tipo', 'Link Roteiro', 'Link Drive', 'Editor', 'Data/Enviado/Cliente', 'Aprov Vídeo']}
    for n, row in enumerate(rd, 2):
        g = lambda k: row[ix[k]].strip() if ix[k] < len(row) else ''
        linhas[norm(g('Sku'))].append(dict(linha=n, tipo=g('Tipo'), roteiro=g('Link Roteiro'), drive=g('Link Drive'),
                                           editor=g('Editor'), ad=g('Data/Enviado/Cliente'), af=g('Aprov Vídeo')))

itens, avisos = [], []
for d, ad, sku, prod, af, nota, audio, obs in ITENS:
    cand = linhas.get(norm(sku), [])
    certa = [c for c in cand if c['ad'] == ad] or cand
    s = certa[-1] if certa else {}
    if not cand:
        avisos.append(f'{sku} {prod[:40]}: SKU não encontrado na aba Setembro/2026')
    elif not [c for c in cand if c['ad'] == ad]:
        avisos.append(f'{sku} {prod[:40]}: nenhuma linha com AD {ad} (usada a linha {s["linha"]}, AD {s["ad"]})')
    ed = {'Bidóia': 'Bidoia'}.get(s.get('editor', ''), s.get('editor', '')) or 'SEM-EDITOR'
    extra = nota
    if s.get('af') == 'Aprovado':
        extra += (' — ' if extra else '') + 'VERIFICAR: na planilha consta Aprovado, mas o cliente pediu esta alteração'
    if ed == 'SEM-EDITOR':
        extra += (' — ' if extra else '') + ('sem editor na planilha' if s else 'SKU não encontrado na aba Setembro/2026')
    itens.append(dict(d=d, sku=sku, prod=prod, af=af, extra=extra, audio=audio, obs=obs, ed=ed, **{k: s.get(k, '') for k in ('tipo', 'roteiro', 'drive')}))


def planilha(path, titulo, sub, cab, larg, rows, link_cols=()):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Ajustes'
    ws['A1'] = titulo
    ws['A1'].font = Font(sz=14, b=True)
    ws['A2'] = sub
    ws['A3'] = f'Pasta-mãe no Drive: {PASTA_MAE}'
    for i, (h, w) in enumerate(zip(cab, larg), 1):
        c = ws.cell(row=5, column=i, value=h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        ws.column_dimensions[c.column_letter].width = w
    for r, vals in enumerate(rows, 6):
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment = WRAP
            if i in link_cols and v:
                c.hyperlink, c.value = v, 'abrir'
                c.font = Font(color='FF0563C1', u='single')
    ws.freeze_panes = 'A6'
    ws.auto_filter.ref = f'A5:{ws.cell(row=5, column=len(cab)).column_letter}{5 + len(rows)}'
    wb.save(path)


por_ed = defaultdict(list)
for x in itens:
    if x['af'] != AU:
        por_ed[x['ed']].append(x)
for ed, xs in sorted(por_ed.items()):
    n_r = sum(x['af'].startswith('REINSERÇÃO') for x in xs)
    planilha(AQUI / f'ajustes_{ed}.xlsx', f'Ajustes do cliente — {ed}',
             f'Diárias {", ".join(sorted({x["d"] for x in xs}))} — {len(xs)} vídeo(s): {n_r} com reinserção de cenas IA, {len(xs) - n_r} de edição',
             ['Diária', 'SKU', 'Produto', 'Tipo', 'O que fazer', 'Observação do cliente', 'Pasta'], [9, 12, 42, 10, 30, 66, 8],
             [[x['d'], x['sku'], x['prod'], x['tipo'], x['af'] + (' — ' + x['extra'] if x['extra'] else ''), x['obs'], x['drive']] for x in xs],
             link_cols=(7,))
    print(f'{ed:12} {len(xs):3}  (reinserção {n_r})')

aud = [x for x in itens if x['audio']]
planilha(AQUI / 'audio-e-roteiro_CONTROLE.xlsx', 'Alterações de áudio e roteiro — pedir regravação',
         f'{len(aud)} vídeo(s). Pronúncia → Drika; corte/troca de trecho → Hector. "VERIFICAR" = conferir se é locução ou lettering.',
         ['Diária', 'SKU', 'Produto', 'O que pedir', 'Editor', 'Observação do cliente', 'Roteiro', 'Pasta'],
         [9, 12, 38, 50, 11, 60, 8, 8],
         [[x['d'], x['sku'], x['prod'], x['audio'], x['ed'], x['obs'], x['roteiro'], x['drive']] for x in aud],
         link_cols=(7, 8))
print('áudio', len(aud))
for a in avisos:
    print('AVISO', a)
